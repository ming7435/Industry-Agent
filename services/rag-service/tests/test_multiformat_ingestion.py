import csv
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from app.chunk import ChunkerConfig, build_chunks
from app.embedding import EmbeddingConfig, embed_chunks
from app.ingestion import parse_document


class FakeEmbeddingClient:
    @property
    def dimension(self) -> int:
        return 2

    def embed_texts(self, texts: list[str]) -> list[list[float]]:
        return [[float(len(text)), 1.0] for text in texts]


class FakeImageDescriber:
    def describe_image(
        self,
        image: bytes,
        *,
        media_type: str,
        filename: str,
        page_number: int,
        prompt: str,
    ) -> str:
        return "主轴装配图，零件编号 A-102，关键尺寸 Φ30。"


class MultiFormatIngestionTests(unittest.TestCase):
    def test_txt_markdown_and_csv_flow_to_vectors(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            txt_path = root / "manual.txt"
            txt_path.write_text("主轴维护步骤：先断电。\n\n检查润滑油。", encoding="utf-8")
            markdown_path = root / "manual.md"
            markdown_path.write_text(
                "# 维护\n\n检查主轴。\n\n| 项目 | 参数 |\n| --- | --- |\n| 油品 | ISO VG 46 |",
                encoding="utf-8",
            )
            csv_path = root / "manual.csv"
            with csv_path.open("w", newline="", encoding="utf-8") as stream:
                csv.writer(stream).writerows([["项目", "参数"], ["油品", "ISO VG 46"]])

            for path, source_format in (
                (txt_path, "txt"),
                (markdown_path, "md"),
                (csv_path, "csv"),
            ):
                with self.subTest(path=path.name):
                    document = parse_document(path)
                    self.assertEqual(document.source_format, source_format)
                    chunks = build_chunks(
                        document,
                        config=ChunkerConfig(min_characters=1),
                    )
                    records = embed_chunks(
                        chunks,
                        FakeEmbeddingClient(),
                        config=EmbeddingConfig(min_characters=1),
                    )
                    self.assertTrue(records)
                    self.assertTrue(all(record.source_format == source_format for record in records))

    def test_docx_and_xlsx_are_parsed_as_structured_tables(self) -> None:
        try:
            from docx import Document
            from openpyxl import Workbook
        except ImportError as exc:
            self.skipTest(f"optional parser dependency unavailable: {exc}")

        with TemporaryDirectory() as directory:
            root = Path(directory)
            docx_path = root / "manual.docx"
            docx = Document()
            docx.add_heading("维护", level=1)
            docx.add_paragraph("检查主轴。")
            table = docx.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "项目"
            table.cell(0, 1).text = "参数"
            table.cell(1, 0).text = "油品"
            table.cell(1, 1).text = "ISO VG 46"
            docx.save(docx_path)

            xlsx_path = root / "manual.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["项目", "参数"])
            sheet.append(["油品", "ISO VG 46"])
            workbook.save(xlsx_path)
            workbook.close()

            docx_result = parse_document(docx_path)
            xlsx_result = parse_document(xlsx_path)
            self.assertTrue(any(block.kind.value == "table" for block in docx_result.blocks))
            self.assertTrue(any(block.kind.value == "table" for block in xlsx_result.blocks))
            self.assertEqual(docx_result.source_format, "docx")
            self.assertEqual(xlsx_result.source_format, "xlsx")

    def test_standalone_image_flows_to_vector_after_vision_recognition(self) -> None:
        try:
            from PIL import Image
        except ImportError as exc:
            self.skipTest(f"optional image dependency unavailable: {exc}")

        with TemporaryDirectory() as directory:
            path = Path(directory) / "machine.png"
            Image.new("RGB", (20, 20), "white").save(path)
            document = parse_document(path, image_describer=FakeImageDescriber())
            chunks = build_chunks(document, config=ChunkerConfig(min_characters=1))
            records = embed_chunks(
                chunks,
                FakeEmbeddingClient(),
                config=EmbeddingConfig(min_characters=1),
            )
            self.assertEqual(document.blocks[0].metadata["recognition_status"], "recognized")
            self.assertTrue(records)
            self.assertTrue(records[0].contains_image)

    def test_dxf_is_parsed_as_cad_with_text_and_geometry_metadata(self) -> None:
        try:
            import ezdxf
        except ImportError as exc:
            self.skipTest(f"optional parser dependency unavailable: {exc}")

        with TemporaryDirectory() as directory:
            path = Path(directory) / "drawing.dxf"
            drawing = ezdxf.new("R2010")
            modelspace = drawing.modelspace()
            modelspace.add_line((0, 0), (100, 0), dxfattribs={"layer": "轮廓线"})
            modelspace.add_text("A-102", dxfattribs={"layer": "零件编号"})
            drawing.saveas(path)

            document = parse_document(path)
            self.assertEqual(document.source_format, "dxf")
            self.assertTrue(document.blocks)
            self.assertTrue(all(block.kind.value == "cad_drawing" for block in document.blocks))
            self.assertEqual(document.metadata["parser"], "ezdxf")
            self.assertIn("A-102", document.blocks[0].content)
            entities = document.metadata["cad_entities"]
            self.assertEqual(len(entities), 2)
            text_entity = next(entity for entity in entities if entity["entity_type"] == "TEXT")
            self.assertEqual(text_entity["layer_name"], "零件编号")
            self.assertEqual(text_entity["text"], "A-102")
            self.assertEqual(text_entity["device_id"], "A-102")
            self.assertTrue(text_entity["handle"])
            self.assertIsNotNone(text_entity["bbox"])
            entity_block = next(block for block in document.blocks if block.metadata.get("source") == "dxf_entity")
            self.assertEqual(entity_block.metadata["entity_handle"], text_entity["handle"])

    def test_dxf_paper_space_entities_are_preserved(self) -> None:
        try:
            import ezdxf
        except ImportError as exc:
            self.skipTest(f"optional parser dependency unavailable: {exc}")

        with TemporaryDirectory() as directory:
            path = Path(directory) / "layout.dxf"
            drawing = ezdxf.new("R2010")
            layout = drawing.layout("Layout1")
            layout.add_text("V-202", dxfattribs={"layer": "图框标注"})
            drawing.saveas(path)

            document = parse_document(path)

            entity = next(
                item for item in document.metadata["cad_entities"]
                if item.get("text") == "V-202"
            )
            self.assertEqual(entity["space_name"], "Layout1")
            self.assertEqual(entity["device_id"], "V-202")

    def test_unsupported_extension_has_clear_error(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "manual.rtf"
            path.write_text("unsupported", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "Unsupported document extension"):
                parse_document(path)


if __name__ == "__main__":
    unittest.main()
